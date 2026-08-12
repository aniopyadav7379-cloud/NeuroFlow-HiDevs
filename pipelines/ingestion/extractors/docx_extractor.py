"""
DOCX extraction via python-docx: paragraphs, table cells, and headers are
extracted as separate ExtractedPage streams, and heading hierarchy is
tracked so the Chunker's `hierarchical` strategy has parent/child sections
to work with (docx with headings -> hierarchical, per chunker.py rules).
"""
import io
import logging

from docx import Document
from docx.table import Table as DocxTable

from pipelines.ingestion.extractors.base import ExtractedPage
from pipelines.ingestion.extractors.markdown import rows_to_markdown_table

logger = logging.getLogger("neuroflow.ingestion.docx")

# python-docx paragraph style names -> heading level, e.g. "Heading 1" -> "h1"
_HEADING_STYLES = {f"Heading {i}": f"h{i}" for i in range(1, 7)}
_HEADING_STYLES["Title"] = "h1"


def has_headings(docx_bytes: bytes) -> bool:
    """Cheap pre-check used by chunker strategy selection: does this
    document have any heading-styled paragraphs at all?"""
    doc = Document(io.BytesIO(docx_bytes))
    return any(p.style.name in _HEADING_STYLES for p in doc.paragraphs if p.style)


def _table_to_page(table: DocxTable, page_number: int, section: str | None) -> ExtractedPage | None:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(c for c in r)]
    if len(rows) < 2:
        return None
    return ExtractedPage(
        page_number=page_number,
        content=rows_to_markdown_table(rows),
        content_type="table",
        metadata={"source": "docx", "section": section},
    )


def extract_docx(docx_bytes: bytes) -> list[ExtractedPage]:
    doc = Document(io.BytesIO(docx_bytes))
    pages: list[ExtractedPage] = []
    page_number = 0
    current_section = None

    # Headers (document header, distinct from heading paragraphs) — one
    # page per section's header text, if non-empty.
    for section in doc.sections:
        header_text = "\n".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
        if header_text:
            page_number += 1
            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    content=header_text,
                    content_type="text",
                    metadata={"source": "docx", "region": "header"},
                )
            )

    # Body: walk paragraphs and tables in document order so table content
    # is attributed to the section it actually appeared under.
    body = doc.element.body
    para_iter = iter(doc.paragraphs)
    table_iter = iter(doc.tables)

    # python-docx doesn't give a single ordered element→object stream out
    # of the box, so we walk body.iterchildren() and match tags back to
    # the parsed paragraph/table objects, preserving encounter order.
    para_by_element = {p._p: p for p in doc.paragraphs}
    table_by_element = {t._tbl: t for t in doc.tables}

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p" and child in para_by_element:
            p = para_by_element[child]
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name if p.style else None
            if style_name in _HEADING_STYLES:
                current_section = text
                page_number += 1
                pages.append(
                    ExtractedPage(
                        page_number=page_number,
                        content=text,
                        content_type="text",
                        metadata={
                            "source": "docx",
                            "level": _HEADING_STYLES[style_name],
                            "section": current_section,
                        },
                    )
                )
            else:
                page_number += 1
                pages.append(
                    ExtractedPage(
                        page_number=page_number,
                        content=text,
                        content_type="text",
                        metadata={"source": "docx", "section": current_section},
                    )
                )
        elif tag == "tbl" and child in table_by_element:
            table_page = _table_to_page(table_by_element[child], page_number + 1, current_section)
            if table_page:
                page_number += 1
                table_page.page_number = page_number
                pages.append(table_page)

    return pages
